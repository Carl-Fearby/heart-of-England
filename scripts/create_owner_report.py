from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "Heart of England Digital Growth and Accessibility Report.docx"
NAVY, INK, MUTED, ROW, LINE = "193B5B", "1E1E1E", "6B6B6B", "F2F4F6", "C9CED3"
CONTENT_WIDTH = 6.94  # Letter width less the 0.78in left and right page margins.

def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr(); el = OxmlElement("w:shd"); el.set(qn("w:fill"), fill); pr.append(el)

def border(cell):
    pr = cell._tc.get_or_add_tcPr(); borders = OxmlElement("w:tcBorders")
    for side in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "5"); el.set(qn("w:color"), LINE); borders.append(el)
    pr.append(borders)

def cell_text(cell, text, bold=False, color=None, size=9.25, font="Georgia"):
    cell.text = ""; p = cell.paragraphs[0]; p.paragraph_format.space_before = Pt(4); p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text); r.bold = bold; r.font.name = font; r.font.size = Pt(size)
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER; border(cell)

def widths(table, values):
    table.autofit = False; table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for cell, value in zip(row.cells, values): cell.width = Inches(value)
    for grid, value in zip(table._tbl.tblGrid.gridCol_lst, values): grid.set(qn("w:w"), str(int(value * 1440)))

def heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}"); p.paragraph_format.space_before = Pt(18 if level == 1 else 12); p.paragraph_format.space_after = Pt(7)
    r = p.add_run(text); r.font.name = "Aptos Display"; r.font.color.rgb = RGBColor.from_string(INK); return p

def para(doc, text, lead=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(8); p.paragraph_format.line_spacing = 1.13
    if lead:
        r = p.add_run(lead + " "); r.bold = True; r.font.name = "Georgia"; r.font.size = Pt(10.5)
    r = p.add_run(text); r.font.name = "Georgia"; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor.from_string(INK); return p

def bullet(doc, text):
    # Avoid Word's theme-dependent list glyphs: some installations substitute
    # them with an empty square. A simple dash is universally available.
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Inches(.18); p.paragraph_format.first_line_indent = Inches(-.18)
    p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.08
    r = p.add_run("- "); r.font.name = "Georgia"; r.font.size = Pt(10.25)
    r = p.add_run(text); r.font.name = "Georgia"; r.font.size = Pt(10.25)

def numbered(doc, text, lead):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(4); p.paragraph_format.line_spacing = 1.08
    r = p.add_run(lead + ". "); r.bold = True; r.font.name = "Georgia"; r.font.size = Pt(10.25)
    r = p.add_run(text); r.font.name = "Georgia"; r.font.size = Pt(10.25)

def table(doc, headers, rows, col_widths):
    # Scale each column plan to the exact body-text measure so the left/right
    # edges of every table align with the surrounding paragraphs.
    scale = CONTENT_WIDTH / sum(col_widths)
    col_widths = [value * scale for value in col_widths]
    t = doc.add_table(rows=1, cols=len(headers))
    for c, text in zip(t.rows[0].cells, headers): cell_text(c, text, True, "FFFFFF", 9.25, "Aptos"); shade(c, NAVY)
    for index, values in enumerate(rows):
        for pos, (c, text) in enumerate(zip(t.add_row().cells, values)):
            cell_text(c, text, pos == 0)
            if index % 2: shade(c, ROW)
    widths(t, col_widths); return t

doc = Document(); sec = doc.sections[0]
sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin = Inches(.72), Inches(.65), Inches(.78), Inches(.78)
styles = doc.styles; styles["Normal"].font.name = "Georgia"; styles["Normal"].font.size = Pt(10.5)
for name, size in (("Heading 1", 21), ("Heading 2", 14), ("Heading 3", 11)):
    styles[name].font.name, styles[name].font.size, styles[name].font.bold = "Aptos Display", Pt(size), True

# Cover treatment follows the established Heart of England report family.
p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(118); p.paragraph_format.space_after = Pt(0)
r = p.add_run("The Heart of England"); r.font.name = "Georgia"; r.font.size = Pt(31); r.font.color.rgb = RGBColor.from_string(NAVY)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(7)
r = p.add_run("Digital Growth and Accessibility Report"); r.bold = True; r.font.name = "Aptos Display"; r.font.size = Pt(25); r.font.color.rgb = RGBColor.from_string(INK)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(13); pr = p._p.get_or_add_pPr(); borders = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "10"); bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), NAVY); borders.append(bottom); pr.append(borders)
p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(22)
r = p.add_run("A practical case for a stronger digital estate and supplier relationship"); r.italic = True; r.font.name = "Georgia"; r.font.size = Pt(13); r.font.color.rgb = RGBColor.from_string(MUTED)
meta = doc.add_table(rows=3, cols=2)
for index, (label, value) in enumerate((("Prepared for", "Stephen Hammon"), ("Scope", "Digital growth, accessibility and website launch readiness"), ("Date", "15 September 2026"))):
    cell_text(meta.rows[index].cells[0], label, True, "FFFFFF", 9.5, "Aptos"); shade(meta.rows[index].cells[0], NAVY)
    cell_text(meta.rows[index].cells[1], value, False, INK, 9.5)
widths(meta, [1.5, 5.44])
heading(doc, "Decision requested", 2)
para(doc, "Approve the Next.js rebuild and appoint a single digital partner to carry the work through launch and into continuous improvement. This gives Heart of England a modern, measurable foundation for website growth, accessibility, search visibility and future campaigns rather than another short-term page refresh.")

doc.add_page_break(); heading(doc, "Executive summary")
para(doc, "This report explains why the replacement website is the stronger long-term option. It compares the legacy WordPress and Elementor experience with the rebuilt Next.js implementation, identifies the issues addressed, and sets out the remaining decisions required before launch.")
heading(doc, "Recommendation", 2)
para(doc, "Proceed with the Next.js rebuild as the future website platform. The work has already moved the site away from a heavily layered WordPress and Elementor implementation, copied the available source content and media locally, and created a cleaner foundation for search, accessibility and future changes. The remaining work is focused on final editorial review and selecting services for forms and consent management.")
heading(doc, "Completed work and why it improves the outcome")
table(doc, ("Completed change", "Why it is a material improvement"), [
    ("Complete navigation restored", "The primary Heart of England navigation now includes conferences, team building, wellbeing workshops, rooms and spaces, events, Christmas parties, Quicken Tree, company information, careers, blog, contact and social links. This protects established discovery routes while giving users and search engines a clearer hierarchy."),
    ("Content and assets migrated locally", "The rebuild has brought across the site structure, editorial pages, room content, Quicken Tree content, menus, room specifications, brochures and locally served imagery. Pages no longer depend on fragile remote upload paths, and known missing legacy files are recorded for replacement rather than appearing as broken content."),
    ("Venue and room journeys rebuilt", "The rooms directory now presents room photography, capacities and clear detail routes. Individual room pages include context, downloadable specifications and layouts explained as room configurations. This turns a dense directory into a usable route from research to enquiry."),
    ("Quicken Tree treated as a first class destination", "The restaurant is retained in the main navigation with its own pages, social routes, booking action and downloadable menus. This gives local diners and event guests a direct path to menus and reservations rather than hiding an important revenue stream behind generic venue content."),
    ("SEO technical foundation added", "Every route has a defined title, description, canonical URL and social-share metadata. The site also exposes robots guidance and a sitemap covering 81 generated routes. These are the technical signals search engines and social platforms need to discover, understand and present the right page."),
    ("Accessible mobile first interface", "The new implementation uses semantic landmarks and headings, a skip link, keyboard-visible focus, responsive navigation, meaningful alternatives for content images and contrast-aware controls. The result is more usable on phones, keyboards and assistive technology and provides a practical WCAG 2.2 AA baseline for ongoing review."),
    ("More resilient media delivery", "A controlled image fallback prevents empty visual frames when an individual source file cannot load. Together with local asset storage, this makes the customer experience more dependable and prevents avoidable loss of confidence at key enquiry points."),
], [2.05, 4.65])
para(doc, "The value is cumulative. Better technical search signals help the correct pages become eligible for discovery; clearer content and mobile-first journeys help visitors understand the offer; accessible interaction helps more people complete the journey; and consistent tracking gives the owner team evidence for where to invest next.")
heading(doc, "Additional technical improvements applied after review")
table(doc, ("Improvement", "Practical impact and justification"), [
    ("Responsive WebP images", "Optimised images now have responsive WebP variants. Browsers can select a smaller file for smaller screens and only request a larger asset when the display needs it. This reduces unnecessary image transfer, which improves the page experience most noticeably on mobile connections."),
    ("Optimised room hero asset", "The Rooms and Spaces hero now uses an optimised WebP source rather than a large JPEG. This is a targeted improvement to an important discovery page where image weight could otherwise delay the first meaningful view."),
    ("Browser caching and baseline security headers", "Long-lived cache rules are in place for optimised images and Next.js static files, while content-type, frame and referrer policies set a safer default delivery posture. Faster repeat visits and a more controlled technical baseline support both customer confidence and operational resilience."),
    ("Structured data expanded", "The site now exposes EventVenue and WebSite schema across the site, Article schema on blog posts and Event schema on the What's on page. This makes the venue, editorial content and scheduled activity easier for search engines to interpret; it improves eligibility for rich presentation but does not guarantee it."),
    ("Metadata and social sharing coverage", "A shared metadata pattern now supplies consistent page titles, descriptions, canonicals, Open Graph and X card information across key content, including contact and blog listing pages. This reduces duplication and makes shared links more likely to present the intended page title, summary and image."),
    ("Indexing hygiene", "The enquiry confirmation page is excluded from search indexing. This prevents a low-value completion page from competing with pages designed to attract new visitors, while still allowing the conversion journey to work normally."),
    ("Static event rendering", "The events listing no longer runs a browser timer for every visitor. Upcoming events are resolved at build time, creating a lighter page. This requires an agreed publishing cadence or automatic scheduled deployment so expired events are removed promptly."),
], [2.05, 4.65])
heading(doc, "A stronger foundation for ongoing digital growth")
para(doc, "The opportunity extends beyond a better-looking website. The proposed platform brings the website, its content, media and customer journeys together in one maintainable foundation, reducing reliance on legacy page-builder blocks, remote media paths and disconnected external services. This makes day-to-day changes more predictable and gives the owner team greater confidence in the customer experience.")
para(doc, "An effective digital partnership joins up strategy, content, technical delivery and measurement. It focuses effort on the pages that generate venue enquiries, restaurant bookings, event discovery and repeat visits; uses evidence to guide improvements; and keeps a clear record of priorities, changes and results. The rebuilt platform is designed to support that way of working.")
heading(doc, "How the rebuild supports commercial growth")
for item in [
    "It gives search engines clear, crawlable pages for rooms, conferences, accommodation, dining, events and seasonal activity, so relevant local and high-intent searches have a stronger route to the right content.",
    "It connects every important journey to an understandable action such as making an enquiry, viewing a room, downloading a menu, checking an event or reserving a table, reducing the chance that a prospective visitor reaches a dead end.",
    "It makes performance measurable. Analytics, Search Console, call and form tracking can show which channels, pages and searches produce qualified interest, allowing budget and content effort to follow evidence rather than assumption.",
    "It creates a durable campaign base. New menus, offers, Christmas activity, weddings, corporate events and Quicken Tree promotions can be launched as focused pages with consistent metadata, imagery and calls to action.",
]: bullet(doc, item)
para(doc, "The website alone cannot guarantee footfall. Local demand, pricing, sales follow-up, event programming and marketing investment all matter. It does, however, remove avoidable digital friction and gives the team the tools to turn more relevant searches and campaign visits into measurable enquiries and bookings.")
heading(doc, "What the legacy site made difficult")
table(doc, ("Area", "Finding from the original site review"), [
    ("Content and layout", "Important venue information was spread across template-specific Elementor blocks. Pages mixed editorial copy, decorative assets, sliders and downloads, which made consistency and maintenance difficult."),
    ("Media management", "Images, PDF menus, room specifications and campaign files depended on remote WordPress upload paths. Some legacy URLs already returned 404 responses."),
    ("Navigation", "The site exposed a large number of services and campaign routes, but the hierarchy was difficult to scan and key actions were distributed across menus and page modules."),
    ("Accessibility", "Legacy visual components relied on layered scripts and image-based interactions. This made keyboard behaviour, focus order, readable text and meaningful content structure harder to control."),
    ("Search and sharing", "Search metadata, canonical URLs, sitemap coverage and social-share previews were inconsistent across the content estate."),
    ("Conversion", "Forms, ticketing, restaurant bookings, virtual tours, vouchers and seasonal campaigns were implemented through multiple third-party embeds and outgoing links."),
], [1.55, 5.15])
heading(doc, "What the replacement has addressed")
for title, body in [
    ("Clean platform", "The project is a focused Next.js application with a small dependency set. It no longer carries the previous database or Drizzle scaffolding."),
    ("Content and media migration", "The migration includes 34 Heart of England pages, 16 Quicken Tree pages, blog content, local room photography, menus, PDF downloads and 671 copied source assets. The asset map covers 777 source URL variants, and source files that now return 404 are recorded rather than silently ignored."),
    ("Venue structure", "The new Rooms and Spaces directory presents the primary event rooms, capacities and routes to individual room detail pages. The Quicken Tree menu page provides six local PDF downloads and a clear restaurant booking action."),
    ("Accessibility baseline", "The rebuild includes semantic page structure, a skip link, keyboard-visible focus styles, responsive navigation, text alternatives for meaningful imagery, readable contrast and a published accessibility statement aligned to WCAG 2.2 AA goals."),
    ("Search foundation", "The build generates a sitemap covering 81 routes, robots metadata, canonical URLs, page titles and descriptions. It also includes Open Graph and X card metadata for social sharing."),
    ("Reliability", "Images are served from the local project rather than relying on the old site. A controlled fallback prevents an empty image frame if a source file is unavailable."),
]: para(doc, body, title + ".")
heading(doc, "Accessibility is a growth and trust issue")
para(doc, "WCAG improvements are not a compliance add-on. Clear headings, keyboard navigation, visible focus, readable contrast, semantic page structure, meaningful image alternatives and responsive layouts make the site easier to use and improve the likelihood that visitors complete an enquiry.")
para(doc, "The new baseline also reduces risk. Accessibility requirements and expectations continue to rise, particularly for public-facing services. A site built with accessible patterns from the start is less expensive to maintain and less likely to need a disruptive retrofit later.")
heading(doc, "Why this is better for the owner")
for item in ["The site can be maintained as a single codebase instead of relying on page-builder modules, old upload paths and embedded scripts.", "The venue, rooms, events and restaurant content can be developed as clear user journeys rather than as copied blocks from older campaigns.", "Search engines and social platforms receive consistent technical signals, including canonical pages, an XML sitemap and share-preview metadata.", "The interface is responsive by default and gives the team a stronger base for WCAG review, content governance and performance optimisation.", "Future content changes such as new events, menus, rooms, campaigns and downloads can be made without rebuilding unrelated page layouts."]: bullet(doc, item)

heading(doc, "Evidence from the completed build")
table(doc, ("Measure", "Current position"), [("Generated routes", "81 static Next.js routes generated successfully"), ("Source content", "50 primary pages across Heart of England and Quicken Tree, plus imported blog posts"), ("Local assets", "671 copied assets and 777 mapped source URL variants"), ("Known unavailable files", "8 source files return 404 from the original websites"), ("Quality checks", "Production build passes; lint has no errors")], [2.15, 4.55])
heading(doc, "Recommended next steps not yet done")
para(doc, "These are owner or deployment decisions from the launch plan. They should be completed before, or immediately after, domain cutover so the new technical foundation is measured, discoverable and operationally complete.")
numbered(doc, "Submit the sitemap to Google Search Console after the production domain is live at heartofengland.uk. This confirms the preferred site to Google and provides the reporting view for indexing, search queries and page performance.", "Google Search Console")
numbered(doc, "Choose a provider and configure measurement and consent before go-live. Analytics must not be enabled without the appropriate consent approach, and the selected platform should support meaningful reporting on enquiries, bookings and campaign outcomes.", "Analytics and cookie consent")
numbered(doc, "Confirm the recipient inbox and handling process for the event-enquiry form in Netlify Forms. This is essential: a high-performing enquiry page only has commercial value when every submission reaches the right team reliably and can be followed up.", "Form notifications")
numbered(doc, "Add a relevant lead image to each blog post. Posts currently share the generic social image; post-specific Open Graph images will make articles more distinctive when shared and can improve the quality of social traffic.", "Blog post images")
numbered(doc, "Set an agreed rebuild and deployment schedule, or automate one, whenever event dates change. Because upcoming events are rendered into static HTML at build time, this prevents past events from remaining visible after their date.", "Rebuild schedule")
heading(doc, "Launch checklist and ownership", 2)
para(doc, "The following checklist turns the outstanding decisions into clear launch controls. Named ownership should be agreed before cutover, with completion evidence retained by the owner team.")
table(doc, ("Workstream", "Action", "Suggested owner", "Completion evidence"), [
    ("Search", "Verify the production domain in Google Search Console and submit the live sitemap.", "Digital partner with owner access", "Verified property, submitted sitemap and an indexing-monitoring date."),
    ("Consent and measurement", "Select an analytics and consent solution, agree the measurement plan, and test consent before analytics tags load.", "Owner decision with digital partner configuration", "Approved provider, documented consent approach and test record."),
    ("Enquiry handling", "Confirm the recipient inbox, fallback recipient and response process for Netlify event-enquiry submissions.", "Heart of England sales team", "Successful test submission and named response owner."),
    ("Content quality", "Review priority pages, downloadable files, post images, contacts and booking links before launch.", "Heart of England content lead with digital partner", "Signed editorial checklist and list of deferred improvements."),
    ("Publishing", "Agree a date-change process and rebuild schedule for events, offers and seasonal activity.", "Digital partner with Heart of England content lead", "Published update process and scheduled deployment responsibility."),
], [1.2, 2.4, 1.4, 1.7])
doc.add_page_break()
heading(doc, "KPI dashboard specification", 2)
para(doc, "The dashboard should be configured after the selected analytics and consent services are in place. Its role is to support commercial decisions, rather than simply report page views.")
table(doc, ("Measure", "Decision it informs", "Source and review cadence"), [
    ("Organic search visibility", "Which venue, room, restaurant and event searches deserve content investment or optimisation.", "Google Search Console, monthly"),
    ("Priority landing-page engagement", "Whether high-value pages help visitors find the appropriate next action.", "Analytics, monthly"),
    ("Enquiry completions and source", "Which channels and pages deliver qualified venue interest.", "Analytics and Netlify Forms, monthly"),
    ("Restaurant booking referrals", "Whether Quicken Tree content and menu journeys contribute to booking interest.", "Analytics and booking-provider reporting, monthly"),
    ("Page experience and accessibility", "Which technical or usability issues should enter the next optimisation cycle.", "Core Web Vitals, accessibility review and support log, quarterly"),
    ("Seasonal campaign performance", "Whether event, Christmas, wedding or promotion pages support the intended enquiry or booking journey.", "Analytics, Search Console and campaign record, after each campaign"),
], [1.7, 2.7, 2.3])
heading(doc, "Ongoing supplier accountability", 2)
para(doc, "After launch, the recommended relationship is a managed improvement cycle rather than ad hoc page changes. It keeps content, technical quality and commercial outcomes connected, with a simple record of priorities, activity and results.")
table(doc, ("Cadence", "Owner value"), [
    ("Monthly reporting", "Visibility of organic search queries, landing pages, enquiries, bookings, calls, campaign performance, page speed and accessibility issues so investment decisions are based on evidence."),
    ("Quarterly optimisation", "A prioritised review of content gaps, conversion journeys, technical SEO, event and seasonal pages, plus WCAG improvements that protect the quality of the public experience."),
    ("Clear ownership", "One partner responsible for coordinating website changes, content publishing, technical maintenance, external-platform links and measurement, with decisions and actions visible to the owner team."),
], [2.05, 4.65])
doc.add_page_break()
heading(doc, "Suggested first 90 days after approval")
para(doc, "This initial programme turns the rebuild into an accountable launch and improvement cycle. It does not assume a particular traffic or revenue outcome; it establishes the measurement, content priorities and review rhythm needed to make informed commercial decisions.")
table(doc, ("Timing", "Priority activity", "Evidence provided to Stephen"), [
    ("Days 1 to 30", "Complete launch decisions: production-domain cutover, Search Console submission, analytics and consent configuration, form-recipient confirmation, and a final check of priority enquiry and booking journeys.", "A launch record showing the live sitemap, tracking configuration, confirmed form routing and any editorial actions still outstanding."),
    ("Days 31 to 60", "Review early search visibility, landing-page behaviour and completed enquiries. Prioritise the first content improvements around high-value venue, room, restaurant, event and seasonal journeys.", "A concise performance baseline and a prioritised improvement backlog, distinguishing evidence, recommendations and work completed."),
    ("Days 61 to 90", "Publish agreed content improvements, add post-specific social imagery, review Core Web Vitals and accessibility findings, and plan the next seasonal or campaign activity.", "A first 90-day review: what changed, what visitors are doing, what has improved, and the recommended plan for the next quarter."),
], [1.15, 3.25, 2.3])
heading(doc, "What success will be measured by", 2)
para(doc, "The proposed reporting should focus on quality rather than headline traffic alone: organic visibility for commercially relevant searches; visits to key venue, rooms, restaurant and event pages; completed enquiry and booking journeys; page experience and accessibility issues; and the contribution of seasonal campaigns. This provides a practical basis for judging the value of ongoing work.")
for section in doc.sections:
    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Heart of England Digital Growth Report   |   "); r.font.name = "Aptos"; r.font.size = Pt(8); r.font.color.rgb = RGBColor.from_string(MUTED)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); footer._p.append(field)
doc.save(OUT)
